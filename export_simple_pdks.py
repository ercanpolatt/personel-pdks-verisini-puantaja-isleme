# -*- coding: utf-8 -*-
"""
===================================================================================================
PDKS Günlük Çalışma Saatleri Sade Raporlayıcı (export_simple_pdks.py)
===================================================================================================
PDKS'de kart basan tüm çalışanları A'dan Z'ye alfabetik sırayla listeler.
1..N günlerinde çalışılan saatleri yazar, çalışılmayan günleri tamamen BOMBOŞ (değersiz) bırakır.
Dinamik takvim desteğiyle tüm ay ve gün sayılarına tam uyumludur.
"""

import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pdks_engine import PDKSEngine

def generate_simple_report(
    output_xlsx="pdks_eylul_gunluk_calisma_saatleri.xlsx",
    output_csv="pdks_eylul_gunluk_calisma_saatleri.csv",
    month=9,
    year=2026
):
    print("1. PDKS ve Puantaj verileri hesaplanıyor...")
    engine = PDKSEngine(pdks_path="pdks.xls", puantaj_path="puantaj.xls", target_month=month, target_year=year)
    engine.load_personnel()
    engine.load_and_process_pdks()
    matrix = engine.get_summary_matrix()

    # Sadece PDKS'de kart basmış olan aktif çalışanları filtrele
    active_matrix = [m for m in matrix if m["total_work_days"] > 0]

    # Türkçe karakter duyarlı alfabetik sıralama
    tr_alphabet = "abcçdefgğhıijklmnoöprsştuüvyz"
    tr_map = {ch: i for i, ch in enumerate(tr_alphabet)}
    def tr_sort_key(name):
        return [tr_map.get(ch, ord(ch) + 100) for ch in str(name).lower()]

    active_sorted = sorted(active_matrix, key=lambda m: tr_sort_key(m["ad_soyad"]))

    print(f"2. Toplam {len(active_sorted)} aktif çalışan tespit edildi. Excel oluşturuluyor...")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{month:02d}_{year}_Saatler"
    ws.views.sheetView[0].showGridLines = True

    # Tema ve Fontlar
    font_title = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    font_hdr   = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_data  = Font(name="Segoe UI", size=9, color="000000")
    font_bold  = Font(name="Segoe UI", size=9, bold=True, color="000000")

    fill_navy     = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_blue_hdr = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_zebra    = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_sunday   = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Pazar günü
    fill_work     = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Çalışılan gün (Açık yeşil)
    fill_total    = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Toplam sütunları

    border_thin = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0")
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left   = Alignment(horizontal="left", vertical="center")
    align_right  = Alignment(horizontal="right", vertical="center")

    tr_months = ["", "Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
    month_name = tr_months[engine.target_month].upper() if 1 <= engine.target_month <= 12 else ""

    # Sütun Başlıkları
    headers = [("Sıra", 6), ("Adı Soyadı", 25), ("Bölüm", 16)]
    for d in range(1, engine.days_in_month + 1):
        day_name = engine.get_day_name_for_day(d, short=True)
        headers.append((f"{d:02d}.{engine.target_month:02d}\n{day_name}", 6))
    headers.append(("Çalışılan\nGün", 9))
    headers.append(("Toplam\nSaat", 10))

    # Başlık Satırı
    ws.merge_cells(f"A1:{get_column_letter(len(headers))}1")
    ws["A1"] = f"FİDE KONSERVE - {month_name} {engine.target_year} PDKS GÜNLÜK ÇALIŞMA SAATLERİ (ALFABETİK LİSTE)"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_navy
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 28

    for c_idx, (h_text, w) in enumerate(headers, start=1):
        col_let = get_column_letter(c_idx)
        cell = ws[f"{col_let}2"]
        cell.value = h_text
        cell.font = font_hdr
        is_sun = False
        if 4 <= c_idx <= 3 + engine.days_in_month:
            d = c_idx - 3
            if engine.is_sunday(d):
                is_sun = True
        cell.fill = fill_sunday if is_sun else fill_blue_hdr
        if is_sun:
            cell.font = Font(name="Segoe UI", size=9, bold=True, color="7F6000")
        cell.alignment = align_center
        cell.border = border_thin
        ws.column_dimensions[col_let].width = w

    # Veri Satırları
    for r_idx, m in enumerate(active_sorted, start=3):
        ws.row_dimensions[r_idx].height = 19
        is_zebra = (r_idx % 2 == 1)
        row_fill = fill_zebra if is_zebra else None

        ws[f"A{r_idx}"] = r_idx - 2
        ws[f"B{r_idx}"] = m["ad_soyad"]
        ws[f"C{r_idx}"] = m["bolum"] if m["bolum"] else ""

        ws[f"A{r_idx}"].alignment = align_center
        ws[f"B{r_idx}"].alignment = align_left
        ws[f"C{r_idx}"].alignment = align_left

        for c in range(1, 4):
            cl = get_column_letter(c)
            ws[f"{cl}{r_idx}"].border = border_thin
            ws[f"{cl}{r_idx}"].font = font_data
            if row_fill:
                ws[f"{cl}{r_idx}"].fill = row_fill

        for d in range(1, engine.days_in_month + 1):
            col_let = get_column_letter(d + 3)
            cell = ws[f"{col_let}{r_idx}"]
            cell.border = border_thin
            cell.alignment = align_center
            h = m["daily_hours"].get(d, 0.0)
            is_sun = engine.is_sunday(d)

            if h > 0:
                cell.value = h
                cell.number_format = "0.0"
                cell.font = font_bold if h > 7.5 else font_data
                cell.fill = fill_work
            else:
                # KULLANICI İSTEĞİ: Boş yerlere çizgi koyma, bomboş kalsın değersiz
                cell.value = None
                if is_sun:
                    cell.fill = fill_sunday
                elif row_fill:
                    cell.fill = row_fill

        # Çalışılan Gün ve Toplam Saat
        col_gun = get_column_letter(3 + engine.days_in_month + 1)
        col_saat = get_column_letter(3 + engine.days_in_month + 2)

        ws[f"{col_gun}{r_idx}"] = m["total_work_days"]
        ws[f"{col_gun}{r_idx}"].alignment = align_center
        ws[f"{col_gun}{r_idx}"].font = font_bold
        ws[f"{col_gun}{r_idx}"].border = border_thin
        ws[f"{col_gun}{r_idx}"].fill = fill_total

        ws[f"{col_saat}{r_idx}"] = m["total_hours"]
        ws[f"{col_saat}{r_idx}"].alignment = align_right
        ws[f"{col_saat}{r_idx}"].font = font_bold
        ws[f"{col_saat}{r_idx}"].number_format = "#,##0.0"
        ws[f"{col_saat}{r_idx}"].border = border_thin
        ws[f"{col_saat}{r_idx}"].fill = fill_total

    saved_excel = output_xlsx
    try:
        wb.save(output_xlsx)
        print(f"3. Excel başarıyla kaydedildi: {output_xlsx}")
    except PermissionError:
        saved_excel = "pdks_eylul_gunluk_calisma_saatleri_guncel.xlsx"
        wb.save(saved_excel)
        print(f"3. BİLGİ: '{output_xlsx}' Excel programında açık olduğu için '{saved_excel}' olarak kaydedildi.")

    # CSV Dışa Aktarımı (Boş hücreler tamamen boş "")
    with open(output_csv, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        for r in range(2, ws.max_row + 1):
            row = []
            for c in range(1, ws.max_column + 1):
                val = ws.cell(r, c).value
                if val is None:
                    val = ""
                elif isinstance(val, str) and "\n" in val:
                    val = val.replace("\n", " ")
                row.append(val)
            writer.writerow(row)
    print(f"4. CSV başarıyla kaydedildi: {output_csv}")

if __name__ == "__main__":
    generate_simple_report()
