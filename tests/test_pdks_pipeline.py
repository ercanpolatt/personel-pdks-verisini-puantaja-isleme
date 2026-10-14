# -*- coding: utf-8 -*-
"""
PDKS Pipeline Test Paketi (test_pdks_pipeline.py)
------------------------------------------------
Vardiya sınıflandırma, mola kuralları, gece çalışması takibi, debounce ve
eksik basım tamamlama mekanizmalarını doğrulayan kapsamlı birim testleri.
"""

import os
from datetime import datetime, date, time, timedelta
import pandas as pd
import pytest
import openpyxl

from pdks_pipeline import (
    PDKSConfig,
    ShiftSchedule,
    DataLoader,
    Cleaner,
    ShiftClassifier,
    HoursCalculator,
    ReportGenerator,
    PDKSPipeline,
    calculate_legal_night_hours,
    parse_flexible_date,
    parse_flexible_time,
    parse_duration_to_minutes
)


# =================================================================================================
# 1. STANDART GÜNDÜZ VARDİYASI MOLA VE ERKEN ÇIKIŞ TESTLERİ
# =================================================================================================

class TestStandardDayShift:
    """A. Standart Gündüz Vardiyası (08:00 – 17:00 / 9-Saatlik Plan) Testleri."""

    def test_standard_day_shift_9h(self):
        """08:00 - 17:00 tam gün çalışma: 9.0s brüt, 1.5s mola, 7.5s net mesai."""
        calc = HoursCalculator()
        brk, net = calc._compute_break_and_net(9.0, "DAY_08_17")
        assert brk == 1.5
        assert net == 7.5

    def test_day_shift_early_departure_8h(self):
        """
        Gündüz vardiyasında 1 saat erken çıkış (8.0 saat bulunma):
        Mola kesintisi 1.5 saat kalır, net çalışma 6.5 saat olur (8.0 - 1.5 = 6.5).
        """
        calc = HoursCalculator()
        brk, net = calc._compute_break_and_net(8.0, "DAY_08_17")
        assert brk == 1.5
        assert net == 6.5

    def test_day_shift_proportional_deduction_5h(self):
        """Gündüz vardiyasında ~5.0 saat kısmi bulunma: 1.0 saat mola düşülür -> Net 4.0 saat."""
        calc = HoursCalculator()
        brk, net = calc._compute_break_and_net(5.0, "DAY_08_17")
        assert brk == 1.0
        assert net == 4.0

    def test_day_shift_proportional_deduction_4h(self):
        """Gündüz vardiyasında ~4.0 saat kısa bulunma: 0.5 saat mola düşülür -> Net 3.5 saat."""
        calc = HoursCalculator()
        brk, net = calc._compute_break_and_net(4.0, "DAY_08_17")
        assert brk == 0.5
        assert net == 3.5

    def test_day_shift_overtime(self):
        """Gündüz vardiyasında fazla mesai (10.0 saat brüt): 1.5 saat mola düşülür -> Net 8.5 saat."""
        calc = HoursCalculator()
        brk, net = calc._compute_break_and_net(10.0, "DAY_08_17")
        assert brk == 1.5
        assert net == 8.5


# =================================================================================================
# 2. DÖNEN 3-VARDİYA VE KRİTİK AYRIM TESTLERİ
# =================================================================================================

class TestRotating3ShiftSystem:
    """B. Dönen 3-Vardiya İşçileri (8 Saatlik Vardiyalar) ve Kritik Ayrım Testleri."""

    def test_rotating_morning_shift_8h(self):
        """08:00 - 16:00 dönen vardiya: 8.0s brüt, 0.5s mola, 7.5s net mesai."""
        calc = HoursCalculator()
        brk, net = calc._compute_break_and_net(8.0, "SHIFT_08_16")
        assert brk == 0.5
        assert net == 7.5

    def test_rotating_evening_shift_8h(self):
        """16:00 - 24:00 dönen akşam vardiyası: 8.0s brüt, 0.5s mola, 7.5s net mesai."""
        calc = HoursCalculator()
        brk, net = calc._compute_break_and_net(8.0, "SHIFT_16_24")
        assert brk == 0.5
        assert net == 7.5

    def test_rotating_night_shift_8h(self):
        """00:00 - 08:00 dönen gece vardiyası: 8.0s brüt, 0.5s mola, 7.5s net mesai."""
        calc = HoursCalculator()
        brk, net = calc._compute_break_and_net(8.0, "SHIFT_00_08")
        assert brk == 0.5
        assert net == 7.5

    def test_critical_distinction_rotating_vs_day_early_departure(self):
        """
        KRİTİK AYRIM TESTİ:
        Aynı 8.0 saatlik brüt bulunma süresine sahip iki personelden:
        1. 08-16 dönen vardiya işçisi -> 7.5 saat net mesai alır.
        2. 08-17 standart gündüz işçisi 1 saat erken çıkmışsa -> 6.5 saat net mesai alır.
        """
        calc = HoursCalculator()
        _, net_rotating = calc._compute_break_and_net(8.0, "SHIFT_08_16")
        _, net_day_early = calc._compute_break_and_net(8.0, "DAY_08_17")

        assert net_rotating == 7.5, "Dönen vardiya 8 saat için 7.5 net saat üretmelidir."
        assert net_day_early == 6.5, "Gündüz vardiyasında erken çıkan işçi için 6.5 net saat üretmelidir."
        assert net_rotating > net_day_early


# =================================================================================================
# 3. YASAL GECE ÇALIŞMASI VE GECE YARISI GEÇİŞLERİ (20:00 - 06:00)
# =================================================================================================

class TestNightWorkAndOvernight:
    """C. Standart Dışı ve Gece Vardiyaları (20:00 - 06:00 Yasal Gece Hesabı)."""

    def test_day_shift_zero_night_hours(self):
        """08:00 - 17:00 gündüz çalışmasında gece saati 0.0 olmalıdır."""
        dt_in = datetime(2026, 9, 2, 8, 0, 0)
        dt_out = datetime(2026, 9, 2, 17, 0, 0)
        night_h = calculate_legal_night_hours(dt_in, dt_out)
        assert night_h == 0.0

    def test_evening_shift_night_hours(self):
        """16:00 - 24:00 akşam vardiyasında 20:00 - 24:00 arası 4.0 saat yasal gece çalışması olmalıdır."""
        dt_in = datetime(2026, 9, 2, 16, 0, 0)
        dt_out = datetime(2026, 9, 3, 0, 0, 0)  # Ertesi gün 00:00
        night_h = calculate_legal_night_hours(dt_in, dt_out)
        assert night_h == 4.0

    def test_night_shift_night_hours(self):
        """00:00 - 08:00 gece vardiyasında 00:00 - 06:00 arası 6.0 saat yasal gece çalışması olmalıdır."""
        dt_in = datetime(2026, 9, 3, 0, 0, 0)
        dt_out = datetime(2026, 9, 3, 8, 0, 0)
        night_h = calculate_legal_night_hours(dt_in, dt_out)
        assert night_h == 6.0

    def test_full_night_shift_20_to_06(self):
        """20:00 - 06:00 gece vardiyasında tam 10.0 saat yasal gece çalışması olmalıdır."""
        dt_in = datetime(2026, 9, 2, 20, 0, 0)
        dt_out = datetime(2026, 9, 3, 6, 0, 0)
        night_h = calculate_legal_night_hours(dt_in, dt_out)
        assert night_h == 10.0

    def test_12h_night_shift_20_to_08(self):
        """20:00 - 08:00 (12 saatlik) gece vardiyasında strictly 20:00-06:00 arası 10.0 saat gece sayılmalıdır."""
        dt_in = datetime(2026, 9, 2, 20, 0, 0)
        dt_out = datetime(2026, 9, 3, 8, 0, 0)
        night_h = calculate_legal_night_hours(dt_in, dt_out)
        assert night_h == 10.0

    def test_evening_overtime_crossing_2000(self):
        """08:00 - 21:30 fazla mesai yapan gündüz işçisinin 20:00-21:30 arası 1.5 saat gece çalışması sayılmalıdır."""
        dt_in = datetime(2026, 9, 2, 8, 0, 0)
        dt_out = datetime(2026, 9, 2, 21, 30, 0)
        night_h = calculate_legal_night_hours(dt_in, dt_out)
        assert night_h == 1.5


# =================================================================================================
# 4. VERİ HİJYENİ: DEBOUNCE, EKSİK BASIM TAMAMLAMA VE SÜRE DENETİMİ
# =================================================================================================

class TestDataHygieneAndCleaners:
    """Veri Kalitesi, Debounce, İki Basım Arası Fark ve Eksik Basım Çözümlemeleri."""

    def test_debounce_drops_rapid_swipes(self):
        """Aynı personelin 5 dakika içindeki mükerrer basımları ayıklanmalıdır."""
        cleaner = Cleaner(PDKSConfig(debounce_threshold_minutes=5.0))
        data = pd.DataFrame([
            {"employee_id": "101", "check_in_date": "02.09.2026", "check_in_time": "07:58:00", "check_out_date": "02.09.2026", "check_out_time": "17:00:00", "shift_code": "08-17"},
            {"employee_id": "101", "check_in_date": "02.09.2026", "check_in_time": "08:01:00", "check_out_date": "02.09.2026", "check_out_time": "17:00:00", "shift_code": "08-17"},  # 3 dk fark -> Drop
            {"employee_id": "101", "check_in_date": "03.09.2026", "check_in_time": "08:00:00", "check_out_date": "03.09.2026", "check_out_time": "17:00:00", "shift_code": "08-17"}
        ])

        cleaned = cleaner.clean_records(data)
        assert len(cleaned) == 2, "5 dakika içindeki mükerrer basım elenmelidir."

    def test_impute_missing_exit(self):
        """Yalnızca giriş basılmışsa çıkış saati 17:00 olarak tamamlanmalı ve AUTO_IMPUTED_EXIT atanmalıdır."""
        cleaner = Cleaner()
        data = pd.DataFrame([
            {"employee_id": "102", "check_in_date": "02.09.2026", "check_in_time": "07:55:00", "check_out_date": None, "check_out_time": None, "shift_code": "08-17"}
        ])

        cleaned = cleaner.clean_records(data)
        assert len(cleaned) == 1
        row = cleaned.iloc[0]
        assert bool(row["imputed_exit"]) is True
        assert row["audit_status"] == "AUTO_IMPUTED_EXIT"
        assert row["cout_time_obj"] == time(17, 0)
        assert row["check_out_dt"] == datetime(2026, 9, 2, 17, 0, 0)

    def test_impute_missing_entry(self):
        """Yalnızca çıkış basılmışsa giriş saati 08:00 olarak tamamlanmalı ve AUTO_IMPUTED_ENTRY atanmalıdır."""
        cleaner = Cleaner()
        data = pd.DataFrame([
            {"employee_id": "103", "check_in_date": None, "check_in_time": None, "check_out_date": "02.09.2026", "check_out_time": "17:05:00", "shift_code": "08-17"}
        ])

        cleaned = cleaner.clean_records(data)
        assert len(cleaned) == 1
        row = cleaned.iloc[0]
        assert bool(row["imputed_entry"]) is True
        assert row["audit_status"] == "AUTO_IMPUTED_ENTRY"
        assert row["cin_time_obj"] == time(8, 0)
        assert row["check_in_dt"] == datetime(2026, 9, 2, 8, 0, 0)

    def test_duration_mismatch_detection(self):
        """Dosyadaki 'Süre' sütunu ile hesaplanan süre arasında > 5 dk fark varsa DURATION_MISMATCH atanmalıdır."""
        cleaner = Cleaner(PDKSConfig(duration_mismatch_threshold_minutes=5.0))
        # 08:00 - 17:00 = 9 saat (540 dk), ama dosyada '08:00' (480 dk) yazıyor (60 dk fark!)
        data = pd.DataFrame([
            {"employee_id": "104", "check_in_date": "02.09.2026", "check_in_time": "08:00:00", "check_out_date": "02.09.2026", "check_out_time": "17:00:00", "raw_duration": "08:00", "shift_code": "08-17"}
        ])

        cleaned = cleaner.clean_records(data)
        row = cleaned.iloc[0]
        assert bool(row["duration_mismatch"]) is True
        assert row["audit_status"] == "DURATION_MISMATCH"
        assert abs(row["duration_delta_min"]) > 5.0


# =================================================================================================
# 5. UÇTAN UCA İŞ AKIŞI VE EXCEL RAPOR ÜRETİMİ (End-to-End Pipeline)
# =================================================================================================

class TestEndToEndPipeline:
    """Tüm motorun uçtan uca çalıştırılması ve Excel sayfalarının doğrulanması."""

    def test_pipeline_on_synthetic_data(self, tmp_path):
        """Sentetik veri üzerinde tam işlem hattını çalıştırır ve 2 sayfalı Excel üretir."""
        sample_data = pd.DataFrame([
            # 1. Normal Gündüz Tam Gün (9h brüt -> 1.5h mola -> 7.5h net)
            {"employee_id": "1", "first_name": "AHMET", "last_name": "YILMAZ", "check_in_date": "02.09.2026", "check_in_time": "08:00:00", "check_out_date": "02.09.2026", "check_out_time": "17:00:00", "shift_code": "08-17", "raw_duration": "09:00"},
            # 2. Gündüz Erken Çıkış 8h (8h brüt -> 1.5h mola -> 6.5h net)
            {"employee_id": "2", "first_name": "MEHMET", "last_name": "DEMİR", "check_in_date": "02.09.2026", "check_in_time": "08:00:00", "check_out_date": "02.09.2026", "check_out_time": "16:00:00", "shift_code": "08-17", "raw_duration": "08:00"},
            # 3. Dönen 3-Vardiya Sabah 8h (8h brüt -> 0.5h mola -> 7.5h net)
            {"employee_id": "3", "first_name": "AYŞE", "last_name": "KAYA", "check_in_date": "02.09.2026", "check_in_time": "08:00:00", "check_out_date": "02.09.2026", "check_out_time": "16:00:00", "shift_code": "08-16", "raw_duration": "08:00"},
            # 4. Eksik Çıkış (Otomatik Impute)
            {"employee_id": "4", "first_name": "FATMA", "last_name": "ÇELİK", "check_in_date": "02.09.2026", "check_in_time": "08:00:00", "check_out_date": None, "check_out_time": None, "shift_code": "08-17"},
            # 5. Gece Vardiyası (16:00 - 24:00 -> 4 saat gece)
            {"employee_id": "5", "first_name": "ALİ", "last_name": "ŞAHİN", "check_in_date": "02.09.2026", "check_in_time": "16:00:00", "check_out_date": "03.09.2026", "check_out_time": "00:00:00", "shift_code": "16-24", "raw_duration": "08:00"}
        ])

        pipeline = PDKSPipeline()
        summary_df, alerts_df = pipeline.process_dataframe(sample_data)

        # Doğrulamalar
        assert len(summary_df) == 5
        # 1. Kayıt
        assert summary_df.iloc[0]["net_hours"] == 7.5
        assert summary_df.iloc[0]["break_hours"] == 1.5
        # 2. Kayıt (Gündüz erken çıkış 8h)
        assert summary_df.iloc[1]["net_hours"] == 6.5
        assert summary_df.iloc[1]["break_hours"] == 1.5
        # 3. Kayıt (3-Vardiya 8h tamamlama)
        assert summary_df.iloc[2]["net_hours"] == 7.5
        assert summary_df.iloc[2]["break_hours"] == 0.5
        # 4. Kayıt (Eksik çıkış)
        assert summary_df.iloc[3]["audit_flag"] == "AUTO_IMPUTED_EXIT"
        assert summary_df.iloc[3]["net_hours"] == 7.5
        # 5. Kayıt (Akşam vardiyası)
        assert summary_df.iloc[4]["night_hours"] == 4.0

        # Alarmlar kontrolü (en az 4. kayıt impute olduğu için alarm listesinde olmalı)
        assert len(alerts_df) >= 1
        assert "4" in alerts_df["employee_id"].values

        # Excel dosya üretimi testi
        out_excel = str(tmp_path / "test_output.xlsx")
        pipeline.reporter.generate_excel_report(summary_df, out_excel)
        assert os.path.exists(out_excel)

        # Excel içerik kontrolü
        wb = openpyxl.load_workbook(out_excel, data_only=True)
        assert "Daily_Summary" in wb.sheetnames
        assert "Audit_Alerts" in wb.sheetnames

        ws1 = wb["Daily_Summary"]
        assert ws1.max_row >= 6  # Başlık + 5 satır + Genel Toplam
        ws2 = wb["Audit_Alerts"]
        assert ws2.max_row >= 3  # Başlık + en az 1 alarm satırı

    def test_pipeline_on_real_pdks_file_if_present(self, tmp_path):
        """Gerçek pdks.xls dosyası varsa uçtan uca çalıştırma testi."""
        if not os.path.exists("pdks.xls"):
            pytest.skip("pdks.xls dosyası bulunamadığı için atlandı.")

        pipeline = PDKSPipeline()
        out_path = str(tmp_path / "real_pdks_report.xlsx")
        res = pipeline.run(input_path="pdks.xls", output_path=out_path)

        assert os.path.exists(res)
        wb = openpyxl.load_workbook(res, data_only=True)
        assert "Daily_Summary" in wb.sheetnames
        assert "Audit_Alerts" in wb.sheetnames
        assert wb["Daily_Summary"].max_row > 100
        assert wb["Audit_Alerts"].max_row > 10
